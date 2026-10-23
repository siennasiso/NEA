"""Build evidence-table.xlsx: Tasks (raw results), SUS (scores) and For write-up (the 4.3 table).

    ../../.venv/bin/python make_xlsx.py            # empty workbook for real results
    ../../.venv/bin/python make_xlsx.py --sample X  # a filled copy at X, only for checking the formulas
"""

import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PINK, HEAD = "C23673", PatternFill("solid", fgColor="FBEAF2")
TASKS = [  # id, short name, screen, criterion (as in the write-up's 4.3 table)
    ("U1", "register and log in", "Register, login", "SC6, SC7"),
    ("U2", "questionnaire", "Questionnaire (10 screens)", "SC8"),
    ("U3", "explain the top match", "Matches", "SC11"),
    ("U4", "staff add, edit, delete", "Staff manage, add animal", "SC12"),
    ("U5", "browse and return", "Browse, profile", "SC4, SC5"),
]
TESTERS = 5
LAST = 1 + TESTERS * len(TASKS)  # last data row on Tasks


def header(ws, row, values):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.font, cell.fill = Font(bold=True, color="7A2049"), HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def build(sample=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    header(ws, 1, ["Tester", "Task", "Completed (Y/N, n.a.)", "Seconds", "Errors (wrong taps)",
                   "Difficulty (1 very easy - 10 very hard)", "Comment"])
    done = DataValidation(type="list", formula1='"Y,N,n.a."', allow_blank=True)
    ws.add_data_validation(done)
    row = 2
    for t in range(1, TESTERS + 1):
        for i, (uid, *_rest) in enumerate(TASKS):
            ws.cell(row, 1, f"Tester {t}")
            ws.cell(row, 2, uid)
            done.add(ws.cell(row, 3))
            if sample:
                staff = t == 5
                if uid == "U4" and not staff:
                    ws.cell(row, 3, "n.a.")
                else:
                    ws.cell(row, 3, "N" if (uid == "U3" and t == 2) else "Y")
                    ws.cell(row, 4, [60, 150, 45, 120, 30][i] + 10 * t)
                    ws.cell(row, 5, (t + i) % 3)
                    ws.cell(row, 6, [2, 3, 6, 5, 2][i] + (t % 2))
            row += 1
    for col, width in zip("ABCDEFG", (10, 7, 14, 9, 12, 16, 50)):
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("SUS")
    header(ws2, 1, ["Tester"] + [f"Q{n}" for n in range(1, 11)] + ["SUS score (0-100)"])
    ws2.cell(1, 13, "Answers 1 (strongly disagree) to 5 (strongly agree). Score = ((each odd Q - 1) + (5 - each even Q)) x 2.5 (Brooke, 1996).")
    for t in range(1, TESTERS + 1):
        r = t + 1
        ws2.cell(r, 1, f"Tester {t}")
        if sample:
            for q in range(1, 11):
                ws2.cell(r, q + 1, (4 + (t % 2)) if q % 2 else (1 + (t % 3 == 0)))
        odd = "+".join(f"({c}{r}-1)" for c in "BDFHJ")
        even = "+".join(f"(5-{c}{r})" for c in "CEGIK")
        ws2.cell(r, 12, f'=IF(COUNT(B{r}:K{r})<10,"",({odd}+{even})*2.5)')
    end = TESTERS + 1
    for r, (label, formula) in enumerate([
        ("Mean", f'=IFERROR(AVERAGE(L2:L{end}),"")'),
        ("Lowest", f'=IFERROR(MIN(L2:L{end}),"")'),
        ("Highest", f'=IFERROR(MAX(L2:L{end}),"")'),
        ("Testers at or above 68", f'=COUNTIF(L2:L{end},">=68")'),
    ], start=end + 2):
        ws2.cell(r, 1, label).font = Font(bold=True)
        ws2.cell(r, 12, formula)
    ws2.column_dimensions["A"].width = 22

    ws3 = wb.create_sheet("For write-up")
    header(ws3, 1, ["Task", "Screen", "Criterion", "Testers", "Completion", "Mean time (s, completed)",
                    "Mean errors", "Mean difficulty (1-10)", "Verdict"])
    rng = lambda c: f"Tasks!${c}$2:${c}${LAST}"  # noqa: E731
    for r, (uid, name, screen, crit) in enumerate(TASKS, start=2):
        ws3.cell(r, 1, f"{uid} {name}")
        ws3.cell(r, 2, screen)
        ws3.cell(r, 3, crit)
        y = f'COUNTIFS({rng("B")},"{uid}",{rng("C")},"Y")'
        n = f'COUNTIFS({rng("B")},"{uid}",{rng("C")},"N")'
        ws3.cell(r, 4, f"={y}+{n}")
        ws3.cell(r, 5, f'=IF(D{r}=0,"",{y}/D{r})').number_format = "0%"
        ws3.cell(r, 6, f'=IFERROR(ROUND(AVERAGEIFS({rng("D")},{rng("B")},"{uid}",{rng("C")},"Y"),0),"")')
        ws3.cell(r, 7, f'=IFERROR(ROUND(AVERAGEIFS({rng("E")},{rng("B")},"{uid}",{rng("C")},"<>n.a."),1),"")')
        ws3.cell(r, 8, f'=IFERROR(ROUND(AVERAGEIFS({rng("F")},{rng("B")},"{uid}",{rng("C")},"<>n.a."),1),"")')
        time_ok = f"AND(F{r}<>\"\",F{r}<=180)" if uid == "U2" else "TRUE"
        ws3.cell(r, 9, f'=IF(D{r}=0,"Not tested",IF(AND(E{r}>=0.8,H{r}<=4,{time_ok}),"Met",'
                       f'IF(E{r}>=0.5,"Partly met","Not met")))')
    s = len(TASKS) + 2
    ws3.cell(s, 1, "SUS (all testers)").font = Font(bold=True)
    ws3.cell(s, 3, "SC15")
    ws3.cell(s, 4, f"=COUNT(SUS!L2:L{TESTERS + 1})")
    ws3.cell(s, 5, "=SUS!L" + str(TESTERS + 3)).number_format = "0.0"
    ws3.cell(s, 9, f'=IF(D{s}=0,"Not tested",IF(E{s}>=68,"Met","Not met"))')
    ws3.cell(s + 2, 1, "Verdict rule: Met = at least 80% completed it and the mean difficulty is 4 or less "
                       "(U2 also needs a mean time of 3 minutes or less, SC8); Partly met = at least 50% completed it; "
                       "otherwise Not met. SC15 is met when the mean SUS score is at least 68 (the average website).")
    ws3.cell(s + 3, 1, "For SUS, column E shows the mean score, not a percentage.")
    for col, width in zip("ABCDEFGHI", (26, 24, 10, 8, 11, 12, 9, 12, 12)):
        ws3.column_dimensions[col].width = width
    return wb


if __name__ == "__main__":
    if "--sample" in sys.argv:
        build(sample=True).save(sys.argv[sys.argv.index("--sample") + 1])
    else:
        build().save("evidence-table.xlsx")
